"""Runnable check for the template.docx override in build_minutes.py: python tests/test_template.py
Builds a bare template (A4, a letterhead header, no bullet/table styles, no numbering part, leftover body text),
then checks the minutes keep the page setup and header, drop the body text, and still get working bullets."""
import json, os, re, subprocess, sys, tempfile, zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

HERE = os.path.dirname(os.path.abspath(__file__))
S = os.path.join(HERE, "..", "plugin", "skills", "meeting-minutes", "scripts")
home = tempfile.mkdtemp()

t = Document()
t.sections[0].page_width, t.sections[0].page_height = Mm(210), Mm(297)
t.sections[0].header.paragraphs[0].text = "LETTERHEAD"
t.add_paragraph("DELETE ME")
raw = os.path.join(home, "raw.docx"); t.save(raw)
with zipfile.ZipFile(raw) as zin, zipfile.ZipFile(os.path.join(home, "template.docx"), "w") as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/numbering.xml":
            continue
        if item.filename == "word/styles.xml":
            data = re.sub(rb'<w:style [^>]*w:styleId="(ListBullet|TableGrid)".*?</w:style>', b"", data, flags=re.S)
        if item.filename in ("word/_rels/document.xml.rels", "[Content_Types].xml"):
            data = re.sub(rb'<(Relationship|Override) [^>]*numbering[^>]*/>', b"", data)
        zout.writestr(item, data)

out = os.path.join(home, "out.docx")
r = subprocess.run([sys.executable, os.path.join(S, "build_minutes.py"), os.path.join(S, "meeting.example.json"), out],
                   env={**os.environ, "MINUTES_HOME": home}, capture_output=True, text=True)
assert r.returncode == 0, r.stderr

d = Document(out)
assert abs(d.sections[0].page_width.mm - 210) < 1, "page size lost"
assert d.sections[0].header.paragraphs[0].text == "LETTERHEAD", "header lost"
assert "DELETE ME" not in "\n".join(p.text for p in d.paragraphs), "template body text kept"
num = d.part.numbering_part.element
bullet_id = d.styles["List Bullet"].element.find(".//" + qn("w:numId")).get(qn("w:val"))
n = next(x for x in num.findall(qn("w:num")) if x.get(qn("w:numId")) == bullet_id)
aid = n.find(qn("w:abstractNumId")).get(qn("w:val"))
assert any(a.get(qn("w:abstractNumId")) == aid for a in num.findall(qn("w:abstractNum"))), "bullet definition missing"
assert "Table Grid" in [s.name for s in d.styles]

os.remove(os.path.join(home, "template.docx"))  # no template: the plain path still works
r = subprocess.run([sys.executable, os.path.join(S, "build_minutes.py"), os.path.join(S, "meeting.example.json"), out],
                   env={**os.environ, "MINUTES_HOME": home}, capture_output=True, text=True)
assert r.returncode == 0, r.stderr
assert Document(out).styles["Normal"].font.name == "Calibri"
print("template override: ok")
