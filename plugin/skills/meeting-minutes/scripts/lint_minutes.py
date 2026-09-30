# -*- coding: utf-8 -*-
"""Mechanical screen of a draft meeting.json against its sources. Works the same for any model or agent.

    python lint_minutes.py meeting.json --transcript talk.whisper.txt --tracker slice.md [--roster roster.local.md]

It catches the classes of error that are cheap to detect and expensive to miss in a compliance record:
  ERROR  dash characters; template placeholders; SPEAKER_xx labels; a hedged time ("approximately"); an action
         owner who is not on the roster
  WARN   a number in the minutes that appears in neither the transcript nor the tracker; a capitalised name that
         appears in neither the roster, the transcript nor the tracker (misheard or invented); an attendee who is not on
         the roster; a sentence with a reason word ("because", "so that", "therefore", "due to" ...) that may be inferred context
  INFO   counts of [to confirm] markers
A clean run does NOT mean the minutes are right. It means the cheap checks passed. Exit code 1 when there are errors.
"""
import argparse
import json
import os
import re
import sys

DATA_DIR = os.environ.get("MINUTES_HOME") or os.path.join(os.path.expanduser("~"), ".claude", "meeting-minutes")

UNITS = {w: i for i, w in enumerate(
    "zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen "
    "seventeen eighteen nineteen".split())}
TENS = {w: 10 * i for i, w in enumerate("_ _ twenty thirty forty fifty sixty seventy eighty ninety".split()) if w != "_"}
ORD = {"first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5, "sixth": 6, "seventh": 7, "eighth": 8,
       "ninth": 9, "tenth": 10, "eleventh": 11, "twelfth": 12, "thirteenth": 13, "fourteenth": 14, "fifteenth": 15,
       "sixteenth": 16, "seventeenth": 17, "eighteenth": 18, "nineteenth": 19, "twentieth": 20, "thirtieth": 30}
STOP = set("""monday tuesday wednesday thursday friday saturday sunday january february march april may june july august
september october november december jan feb mar apr jun jul aug sep sept oct nov dec mon tue tues wed thu thur thurs fri sat sun schedule appendix teams microsoft word
excel refer note noted pending approved for noting the and with from into that this these those there their they them
""".split())
REASON_RE = re.compile(r"\b(because|so that|in order to|therefore|as a result|which means|due to|leading to|hence|thus)\b", re.I)
HEDGE_RE = re.compile(r"\b(estimated|approximately|approx|about|around|roughly)\b|~", re.I)
NUM_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")
CAP_RE = re.compile(r"\b[A-Z][A-Za-z'\-]{2,}\b")


def word_numbers(text):
    """Integers spelled out in words ('twelve', 'two thousand', 'fifth')."""
    found, cur, total, seen = set(), 0, 0, False
    for tok in re.findall(r"[a-z]+", text.lower()) + ["."]:
        if tok in UNITS:
            cur += UNITS[tok]; seen = True
        elif tok in TENS:
            cur += TENS[tok]; seen = True
        elif tok in ORD:
            found.add(ORD[tok]); cur += ORD[tok]; seen = True
        elif tok == "hundred":
            cur = max(cur, 1) * 100; seen = True
        elif tok == "thousand":
            total += max(cur, 1) * 1000; cur = 0; seen = True
        elif tok == "and" and seen:
            continue
        else:
            if seen:
                found.add(total + cur)
                if cur:
                    found.add(cur)
            cur, total, seen = 0, 0, False
    return found


def digit_numbers(text):
    out = set()
    for m in NUM_RE.finditer(text):
        s = m.group(0).replace(",", "")
        try:
            f = float(s)
        except ValueError:
            continue
        out.add(int(f) if f == int(f) else f)
        if re.match(r"\d", text[m.end():m.end() + 1] or "") is None and text[m.end():m.end() + 1].lower() == "k":
            out.add(int(f * 1000))
    return out


def load_text(path):
    if not path:
        return ""
    if path.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        return "\n".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    with open(path, encoding="utf-8-sig") as f:
        return f.read()


def strings_of(m):
    """(where, text) for every free-text field that is checked against the sources."""
    for i, it in enumerate(m.get("items", []), 1):
        if "section" in it:
            continue
        yield f"item '{it.get('title', '?')[:40]}' title", it.get("title", "")
        for b in it.get("discussion", []):
            yield f"item '{it.get('title', '?')[:40]}' discussion", b
        for b in it.get("actions", []):
            yield f"item '{it.get('title', '?')[:40]}' action", b
    for s in m.get("schedules", []):
        for row in s.get("rows", []):
            yield f"{s.get('title', 'schedule')[:30]} row", " ".join(str(x) for x in row)
        if s.get("note"):
            yield f"{s.get('title', 'schedule')[:30]} note", s["note"]
    for n in m.get("notes", []):
        yield "notes", n


def all_strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for v in obj.values():
            yield from all_strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from all_strings(v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("meeting_json")
    ap.add_argument("--transcript", help="the speaker-labelled transcript (.txt)")
    ap.add_argument("--tracker", help="the tracker slice (.md) or the tracker (.xlsx)")
    ap.add_argument("--roster", default=os.path.join(DATA_DIR, "roster.local.md"))
    ap.add_argument("--allow", default="", help="extra words to accept, comma-separated")
    a = ap.parse_args()

    with open(a.meeting_json, encoding="utf-8-sig") as f:
        m = json.load(f)
    transcript, tracker = load_text(a.transcript), load_text(a.tracker)
    roster = load_text(a.roster) if a.roster and os.path.exists(a.roster) else ""
    cfg = {}
    cp = os.path.join(DATA_DIR, "config.json")
    if os.path.exists(cp):
        with open(cp, encoding="utf-8-sig") as f:
            cfg = json.load(f)
    if not transcript and not tracker:
        sys.exit("give at least one of --transcript / --tracker")

    src_nums = set()
    for t in (transcript, tracker):
        src_nums |= digit_numbers(t) | word_numbers(t)
    known = set(w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'\-]+", " ".join([transcript, tracker, roster, " ".join(map(str, cfg.values()))])))
    known |= STOP | {w.strip().lower() for w in a.allow.split(",") if w.strip()}
    roster_tokens = set(w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'\-]+", roster))

    errors, warns, info = [], [], []

    # ---- ERROR: dashes, placeholders, speaker labels, hedged time
    for s in all_strings(m):
        if "—" in s or "–" in s:
            errors.append(f"dash character in: {s[:80]!r} (write ranges as 'to', or rephrase)")
        if re.search(r"\[(?!to confirm\]|reason unclear\])[^\]]{2,}\]", s):
            errors.append(f"template placeholder left in: {s[:80]!r}")
        if re.search(r"SPEAKER_\d+", s):
            errors.append(f"unnamed speaker label in: {s[:80]!r}")
    if HEDGE_RE.search(str(m.get("time", ""))):
        errors.append(f"time is hedged: {m.get('time')!r} (derive it or ask; no 'estimated'/'about' in a compliance header)")
    for k in ("date", "time", "venue", "closing"):
        if not str(m.get(k, "")).strip():
            errors.append(f"missing field: {k}")

    # ---- owners and attendees against the roster
    if roster:
        for it in m.get("items", []):
            for act in it.get("actions", []) if "section" not in it else []:
                mm = re.match(r"^([^:]{1,45}):\s", act)
                if not mm:
                    continue
                for tok in re.findall(r"[A-Za-z][A-Za-z'\-]+", mm.group(1)):
                    if tok.lower() not in roster_tokens and tok.lower() not in ("and", "all", "team", "everyone"):
                        errors.append(f"action owner {mm.group(1)!r} is not on the roster ({tok!r}): {act[:70]!r}")
                        break
        for line in list(m.get("attendance", [])) + list(m.get("not_present", [])):
            name = line.split(",")[0]
            bad = [t for t in re.findall(r"[A-Za-z][A-Za-z'\-]+", name) if t.lower() not in roster_tokens]
            if bad:
                warns.append(f"attendee not on the roster ({', '.join(bad)}): {line[:70]!r}")
    else:
        info.append("no roster found: owner and attendee names were not checked")

    # ---- numbers and names against the sources
    seen_nums, seen_names = set(), set()
    for where, text in strings_of(m):
        scan = re.sub(r"Schedule\s+\d+", "Schedule", text)
        for n in digit_numbers(scan):
            if n not in src_nums and (where, n) not in seen_nums:
                seen_nums.add((where, n))
                warns.append(f"number {n} not found in the transcript or tracker ({where}); fine only if you computed it (e.g. a date from a weekday): {text[:80]!r}")
        for sent in re.split(r"(?<=[.!?])\s+|\s+-\s+|:\s+", text):
            words = sent.split()
            for w in CAP_RE.findall(" ".join(words[1:])):   # the first word of a sentence is capitalised anyway
                lw = re.sub(r"'s$", "", w.lower().strip("'"))      # "Timothy's" -> "timothy"
                parts = [p for p in lw.split("-") if p]             # "China-specific" is fine if every part is known
                if not all(p in known for p in parts) and lw not in known and lw not in seen_names:
                    seen_names.add(lw)
                    warns.append(f"name/term {w!r} not found in the roster, transcript or tracker ({where}): check the spelling")
        for r in REASON_RE.finditer(text):
            warns.append(f"reason word {r.group(0)!r}: is this stated in the source, or inferred? ({where}): {text[:90]!r}")

    n_conf = sum(s.count("[to confirm]") + s.count("[reason unclear]") for s in all_strings(m))
    if n_conf:
        info.append(f"{n_conf} [to confirm]/[reason unclear] marker(s): make sure each is on the review list for the user")

    for label, rows in (("ERROR", errors), ("WARN", warns), ("INFO", info)):
        for r in rows:
            print(f"{label}: {r}")
    print(f"\n{len(errors)} error(s), {len(warns)} warning(s). A clean run only means the cheap checks passed; it does not prove the minutes are right.")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
