# -*- coding: utf-8 -*-
"""Final pass on a content-FROZEN, hand-edited minutes docx. Not a rebuild.

    python sign_minutes.py minutes.docx --date 7/9/2026 [--signature path.png]

  - cantSplit the appendix schedule rows so they never wrap across a page
  - fill the Prepared-by date + the scribe's signature image

Pagination of the main discussion table is handled at build time in build_minutes.py
(keepLines + keepNext). Do NOT reintroduce continuation-row surgery here; see
references/format.md, "Pagination".
"""
import argparse
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

DATA_DIR = os.environ.get('MINUTES_HOME') or os.path.join(os.path.expanduser('~'), '.claude', 'meeting-minutes')


def cantsplit(tr):
    trPr = tr.find(qn('w:trPr'))
    if trPr is None:
        trPr = OxmlElement('w:trPr'); tr.insert(0, trPr)
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('--date', required=True, help='signing date as written on the form, e.g. 7/9/2026')
    ap.add_argument('--signature', default=os.path.join(DATA_DIR, 'signature.png'))
    a = ap.parse_args()
    if not os.path.exists(a.signature):
        raise SystemExit('signature image not found: %s (run the setup wizard)' % a.signature)

    doc = Document(a.docx)

    # find the sign-off table by content, not by index
    sign_idx = next((i for i, t in enumerate(doc.tables)
                     if t.rows and t.rows[0].cells[0].text.strip().startswith('Prepared by')), None)
    if sign_idx is None:
        raise SystemExit("no 'Prepared by' sign-off table found")

    # appendix schedules = every table after the sign-off table
    for t in doc.tables[sign_idx + 1:]:
        for tr in t._tbl.findall(qn('w:tr')):
            cantsplit(tr)

    prep = doc.tables[sign_idx].rows[2].cells[0]
    p = prep.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run('Date: %s        Signature: ' % a.date)
    run.font.size = Pt(10.5)
    p.add_run().add_picture(a.signature, width=Inches(1.1))

    doc.save(a.docx)
    print('SIGNED', a.docx)


if __name__ == '__main__':
    main()
