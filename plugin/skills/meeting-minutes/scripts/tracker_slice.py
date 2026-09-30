# -*- coding: utf-8 -*-
"""Cut a weekly tracker (.xlsx) down to the ONE week the meeting reports on.

    python tracker_slice.py tracker.xlsx --meeting-date 2026-09-22 [-o slice.md]
    python tracker_slice.py tracker.xlsx --week "14 Sep" [-o slice.md]
    python tracker_slice.py tracker.xlsx --list-weeks

A meeting reports on the week BEFORE it. The meeting-week column and later hold planned items, not a record, so
they are deliberately left out. With --meeting-date the reported week is the one starting on the Monday before the
Monday of the meeting's week (22 Sep 2026, a Tuesday, reports the week of Mon 14 Sep).

Trackers are usually one sheet per person, and some lay weeks out as ROWS ("WEEK OF 14 SEP" in column A) while others
use COLUMNS (week labels in row 1). Both are handled. Labels have no year, so the meeting's year is assumed.
Anything not recognised is listed under "skipped" so a human can look. Output is Markdown, UTF-8.
"""
import argparse
import datetime as dt
import re
import sys

import openpyxl

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
LABEL_RE = re.compile(r"week\s*of\s*(\d{1,2})\s*(?:st|nd|rd|th)?\s*([a-z]{3,9})", re.I)


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d %b %Y")
    return re.sub(r"\s+", " ", str(v)).strip()


def label_date(cell, year):
    """(day, month) of a 'WEEK OF 14 SEP' style label, or a real date cell; else None."""
    if isinstance(cell, (dt.datetime, dt.date)):
        return dt.date(cell.year, cell.month, cell.day)
    if isinstance(cell, str):
        m = LABEL_RE.search(cell)
        if m and m.group(2)[:3].lower() in MONTHS:
            try:
                return dt.date(year, MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
            except ValueError:
                return None
    return None


def monday(d):
    return d - dt.timedelta(days=d.weekday())


def find_weeks(ws, year):
    """[(orientation, index, date, label)] for every week label on the sheet."""
    out = []
    for c in range(1, min(ws.max_column, 200) + 1):
        d = label_date(ws.cell(1, c).value, year)
        if d:
            out.append(("col", c, d, clean(ws.cell(1, c).value)))
    for r in range(2, min(ws.max_row, 400) + 1):   # row 1 was scanned above (a label in A1 counts once, as a column)
        d = label_date(ws.cell(r, 1).value, year)
        if d:
            out.append(("row", r, d, clean(ws.cell(r, 1).value)))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--meeting-date", help="YYYY-MM-DD (the meeting reports on the week before)")
    ap.add_argument("--week", help='the week to extract, e.g. "14 Sep" (Monday date as labelled)')
    ap.add_argument("--year", type=int, help="year of the labels (default: the meeting date's year, else this year)")
    ap.add_argument("--list-weeks", action="store_true")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()

    md = dt.date.fromisoformat(a.meeting_date) if a.meeting_date else None
    year = a.year or (md.year if md else dt.date.today().year)
    target = None
    if a.week:
        m = re.match(r"\s*(\d{1,2})\s*([A-Za-z]{3,9})", a.week)
        if not m or m.group(2)[:3].lower() not in MONTHS:
            sys.exit('--week must look like "14 Sep"')
        target = dt.date(year, MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
    elif md:
        target = monday(md) - dt.timedelta(days=7)
    elif not a.list_weeks:
        sys.exit("give --meeting-date, --week, or --list-weeks")

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    lines = []
    skipped = []
    if target:
        lines.append(f"# Tracker slice: week of {target.strftime('%d %b %Y')} (the week the meeting reports on)")
        lines.append(f"Source: {a.xlsx}. Only this week is shown; the meeting week and later are planned items, not a record.\n")
    for ws in wb.worksheets:
        weeks = find_weeks(ws, year)
        if a.list_weeks:
            lines.append(f"## {ws.title}: " + (", ".join(f"{w[3]} ({w[0]} {w[1]})" for w in weeks) or "no week labels found"))
            continue
        if not weeks:
            skipped.append(f"{ws.title} (no 'WEEK OF ...' labels found; read it by hand if it matters)")
            continue
        # match the Monday, tolerating a label dated 1 day either side (Sunday/Tuesday-dated weeks)
        hits = [w for w in weeks if abs((w[2] - target).days) <= 1]
        if not hits:
            skipped.append(f"{ws.title} (has weeks {weeks[0][3]} .. {weeks[-1][3]} but none for {target.strftime('%d %b')})")
            continue
        lines.append(f"## Sheet: {ws.title}")
        for orient, idx, d, label in hits:
            if len(hits) > 1:
                lines.append(f"(label '{label}' appears more than once on this sheet: all copies are shown, check which is right)")
            if orient == "col":
                lines.append(f"### {label} (column {idx})")
                for r in range(2, ws.max_row + 1):
                    val = clean(ws.cell(r, idx).value)
                    name = clean(ws.cell(r, 1).value)
                    if val and val.upper() not in ("N/A", "NA", "-"):
                        lines.append(f"- **{name or f'row {r}'}**: {val}")
                    elif val:
                        lines.append(f"- {name or f'row {r}'}: {val}")
            else:
                lines.append(f"### {label} (row {idx})")
                # column titles sit in the rows above the first week row
                first = min(w[1] for w in weeks if w[0] == "row")
                for c in range(2, ws.max_column + 1):
                    val = clean(ws.cell(idx, c).value)
                    if not val:
                        continue
                    title = " / ".join(t for t in (clean(ws.cell(r, c).value) for r in range(1, first)) if t)
                    lines.append(f"- **{title or f'column {c}'}**: {val}")
        lines.append("")
    if skipped and not a.list_weeks:
        lines.append("## Skipped (no matching week found)")
        lines += [f"- {s}" for s in skipped]
    text = "\n".join(lines) + "\n"
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"wrote {a.out} ({len(text)} chars)")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
